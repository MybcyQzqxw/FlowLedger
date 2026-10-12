"""Workbook look & feel: 楷体 for Chinese, Times New Roman for Western text, total row."""
from __future__ import annotations

from datetime import datetime
from typing import List, Tuple

from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import AMOUNT_FORMAT, COLUMN_WIDTHS, HEADERS, TIME_FORMAT, TOTAL_LABEL

FONT_CN = "楷体"
FONT_EN = "Times New Roman"
BODY_SIZE = 11

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
ZEBRA_FILL = PatternFill("solid", fgColor="EEF4FB")
TOTAL_FILL = PatternFill("solid", fgColor="FFE699")

_THIN = Side(style="thin", color="BCCADB")
BODY_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GOLD = Side(style="medium", color="BF8F00")
TOTAL_BORDER = Border(left=_THIN, right=_THIN, top=Side(style="double", color="BF8F00"), bottom=_GOLD)

# 0-based column index -> horizontal alignment
_ALIGN = ("center", "left", "center", "center", "center", "left", "right", "left")
_WRAP_COLUMNS = {5, 7}
_MAX_WIDTH = 60


def _is_cn(ch: str) -> bool:
    # Latin/Greek/digits/ASCII symbols are "Western"; CJK, full-width forms,
    # curly quotes, dashes etc. are all >= U+2000 and go with the Chinese font.
    return ord(ch) >= 0x2000


def display_width(text: str) -> int:
    return sum(2 if _is_cn(c) else 1 for c in text)


def _split_runs(text: str) -> List[Tuple[bool, str]]:
    runs: List[Tuple[bool, str]] = []
    for ch in text:
        cn = _is_cn(ch) if not ch.isspace() else (runs[-1][0] if runs else False)
        if runs and runs[-1][0] == cn:
            runs[-1] = (cn, runs[-1][1] + ch)
        else:
            runs.append((cn, ch))
    return runs


def set_text(cell, text: str, size: float = BODY_SIZE, bold: bool = False, color: str = "000000") -> None:
    """Write text so Chinese characters use 楷体 and everything else Times New Roman."""
    has_cn = any(_is_cn(c) for c in text)
    has_en = any(not _is_cn(c) and not c.isspace() for c in text)
    cell.font = Font(name=FONT_CN if has_cn and not has_en else FONT_EN, size=size, bold=bold, color=color)
    if has_cn and has_en:
        cell.value = CellRichText(
            *[
                TextBlock(
                    InlineFont(rFont=FONT_CN if cn else FONT_EN, sz=size, b=bold, color=color),
                    chunk,
                )
                for cn, chunk in _split_runs(text)
            ]
        )
    else:
        cell.value = text
        cell.data_type = "s"  # never let a remark starting with "=" become a formula


def set_western(cell, value, size: float = BODY_SIZE, bold: bool = False, color: str = "000000") -> None:
    cell.value = value
    cell.font = Font(name=FONT_EN, size=size, bold=bold, color=color)


def style_header(ws) -> None:
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        set_text(cell, title, size=12, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BODY_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def write_records(ws, records) -> None:
    """Write data rows + the total row (assumes rows 2.. are already cleared)."""
    widths = [max(COLUMN_WIDTHS[i], display_width(HEADERS[i]) + 6) for i in range(len(HEADERS))]
    wrap_overflow_rows = set()

    for r, rec in enumerate(records, start=2):
        for c, value in enumerate(rec.to_row()):
            cell = ws.cell(row=r, column=c + 1)
            if isinstance(value, datetime):
                set_western(cell, value)
                cell.number_format = TIME_FORMAT
                width = 21
            elif c == 6:
                set_western(cell, value)
                if value is not None:
                    cell.number_format = AMOUNT_FORMAT
                width = len(f"{value or 0:,.2f}") + 4
            else:
                text = "" if value is None else str(value)
                set_text(cell, text)
                if c == 2:
                    cell.number_format = "@"
                width = display_width(text) + 3
            widths[c] = min(max(widths[c], width), _MAX_WIDTH)
            if c in _WRAP_COLUMNS and width > _MAX_WIDTH:
                wrap_overflow_rows.add(r)
            cell.alignment = Alignment(
                horizontal=_ALIGN[c],
                vertical="center",
                wrap_text=c in _WRAP_COLUMNS,
                indent=0 if _ALIGN[c] == "center" else 1,
            )
            cell.border = BODY_BORDER
            if r % 2 == 1:
                cell.fill = ZEBRA_FILL
        ws.row_dimensions[r].height = None if r in wrap_overflow_rows else 22

    last = len(records) + 1
    if records:
        _write_total(ws, last)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    # the total is rendered larger than body text, give its column enough room
    amount_col = ws.column_dimensions[get_column_letter(7)]
    amount_col.width = max(amount_col.width or 0, 20)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(last, 1)}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _write_total(ws, last_data_row: int) -> None:
    row = last_data_row + 1
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.alignment = Alignment(horizontal="right", vertical="center")
    label = ws.cell(row=row, column=6)
    set_text(label, TOTAL_LABEL, size=14, bold=True, color="7F3F00")
    total = ws.cell(row=row, column=7)
    set_western(total, f"=SUM(G2:G{last_data_row})", size=15, bold=True, color="C00000")
    total.number_format = AMOUNT_FORMAT
    ws.row_dimensions[row].height = 32
