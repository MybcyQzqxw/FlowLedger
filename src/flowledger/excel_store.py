"""Excel persistence: dedupe, chronological ordering and safe saving."""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import List, Optional, Union

from openpyxl import Workbook, load_workbook

from .constants import HEADERS, TOTAL_LABEL
from .excel_style import style_header, write_records


class ExcelError(Exception):
    """Raised for any problem reading or writing the target workbook."""


_TIME_RE = re.compile(
    r"(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})\s*日?\s*(\d{1,2})\s*[:：时]\s*(\d{1,2})(?:\s*[:：分]\s*(\d{1,2}))?"
)


def parse_datetime(text: str) -> Optional[datetime]:
    m = _TIME_RE.search(text)
    if not m:
        return None
    y, mo, d, h, mi, s = m.groups()
    try:
        return datetime(int(y), int(mo), int(d), int(h), int(mi), int(s or 0))
    except ValueError:
        return None


def parse_amount(text: str) -> Optional[float]:
    cleaned = re.sub(r"[^\d.,]", "", text)
    m = re.fullmatch(r"(\d[\d.,]*?)[.,](\d{2})", cleaned)
    if m:
        integer = re.sub(r"[.,]", "", m.group(1))
        return float(f"{integer}.{m.group(2)}")
    if cleaned.isdigit():
        return float(cleaned)
    return None


def _norm_text(value: str) -> str:
    return re.sub(r"[\W_]+", "", value or "").lower()


def _norm_account(value: str) -> str:
    return re.sub(r"[\s＊•●·xX×]", lambda m: "*" if not m.group(0).isspace() else "", value or "")


@dataclass
class Record:
    time: Union[datetime, str, None]
    payer_name: str
    payer_account: str
    tx_type: str
    channel: str
    detail: str
    amount: Optional[float]
    remark: str

    def is_duplicate_of(self, other: "Record") -> bool:
        """Same transaction? Tolerant to OCR noise in the free-text fields.

        Time, amount and payer account must match exactly; the remark must match
        after dropping spaces/punctuation (or be very similar)."""
        if self._time_key() != other._time_key():
            return False
        if self.amount is None or other.amount is None:
            if self.amount != other.amount:
                return False
        elif round(self.amount * 100) != round(other.amount * 100):
            return False
        if _norm_account(self.payer_account) != _norm_account(other.payer_account):
            return False
        a, b = _norm_text(self.remark), _norm_text(other.remark)
        if a == b:
            return True
        if not a or not b:
            return False
        return SequenceMatcher(None, a, b).ratio() >= 0.8

    def _time_key(self) -> str:
        if isinstance(self.time, datetime):
            return self.time.isoformat()
        return str(self.time or "").strip()

    def sort_key(self) -> tuple:
        if isinstance(self.time, datetime):
            return (0, self.time)
        return (1, datetime.max)

    def to_row(self) -> list:
        return [
            self.time,
            self.payer_name,
            self.payer_account,
            self.tx_type,
            self.channel,
            self.detail,
            self.amount,
            self.remark,
        ]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _row_to_record(values) -> Optional[Record]:
    values = list(values) + [None] * (len(HEADERS) - len(values))
    if all(_text(v) == "" for v in values[: len(HEADERS)]):
        return None
    if _text(values[0]) == "" and _text(values[5]) == TOTAL_LABEL:
        return None  # the auto-generated total row
    raw_time = values[0]
    if isinstance(raw_time, datetime):
        time: Union[datetime, str, None] = raw_time.replace(microsecond=0)
    elif _text(raw_time) == "":
        time = None
    else:
        time = parse_datetime(_text(raw_time)) or _text(raw_time)
    raw_amount = values[6]
    if isinstance(raw_amount, (int, float)):
        amount: Optional[float] = float(raw_amount)
    else:
        amount = parse_amount(_text(raw_amount)) if _text(raw_amount) else None
    return Record(
        time=time,
        payer_name=_text(values[1]),
        payer_account=_text(values[2]),
        tx_type=_text(values[3]),
        channel=_text(values[4]),
        detail=_text(values[5]),
        amount=amount,
        remark=_text(values[7]),
    )


class ExcelStore:
    def __init__(self, path: Union[str, Path]):
        self.path = Path(path)

    # -- public API ---------------------------------------------------
    def check_writable(self) -> None:
        """Fail early (before OCR work) if the workbook cannot be written."""
        parent = self.path.parent
        if not parent.is_dir():
            raise ExcelError(f"Excel 所在目录不存在：{parent}")
        if self.path.exists():
            try:
                with open(self.path, "r+b"):
                    pass
            except PermissionError:
                raise ExcelError("Excel 文件无法写入，可能正在被 Excel 打开，请先关闭后重试")
            except OSError as exc:
                raise ExcelError(f"无法访问 Excel 文件：{exc}")

    def find_duplicate(self, record: Record) -> Optional[int]:
        """Excel row number of an identical existing record, or None."""
        if not self.path.exists():
            return None
        _, ws = self._open()
        for row_no, values in enumerate(
            ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True), start=2
        ):
            existing = _row_to_record(values)
            if existing is not None and existing.is_duplicate_of(record):
                return row_no
        return None

    def add(self, record: Record) -> Optional[int]:
        """Insert `record` keeping chronological order.

        Returns the Excel row number it ended up in, or None for a duplicate."""
        wb, ws = self._open()
        records = self._read(ws)
        if any(r.is_duplicate_of(record) for r in records):
            return None
        records.append(record)
        records.sort(key=Record.sort_key)  # stable: ties keep insertion order
        self._rewrite(ws, records)
        self._save(wb)
        return next(i for i, r in enumerate(records, start=2) if r is record)

    # -- internals ------------------------------------------------------
    def _open(self):
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "进账明细"
            return wb, ws
        try:
            wb = load_workbook(self.path)
        except PermissionError:
            raise ExcelError("Excel 文件无法读取，可能被占用")
        except Exception as exc:
            raise ExcelError(f"无法打开 Excel 文件（需要有效的 .xlsx）：{exc}")
        ws = wb.active
        header = [_text(ws.cell(row=1, column=c + 1).value) for c in range(len(HEADERS))]
        if any(header) and tuple(header) != HEADERS:
            raise ExcelError(
                "Excel 表头与预期不一致，期望：" + "、".join(HEADERS) + "；实际：" + "、".join(header)
            )
        return wb, ws

    @staticmethod
    def _read(ws) -> List[Record]:
        records = []
        for values in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
            rec = _row_to_record(values)
            if rec is not None:
                records.append(rec)
        return records

    @staticmethod
    def _rewrite(ws, records: List[Record]) -> None:
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        style_header(ws)
        write_records(ws, records)
        ws.parent.calculation.fullCalcOnLoad = True  # make Excel compute the total on open

    def _save(self, wb) -> None:
        tmp = self.path.with_name(f".{self.path.name}.tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError:
            raise ExcelError("写入 Excel 失败，文件可能正在被 Excel 打开，请先关闭后重试")
        except OSError as exc:
            raise ExcelError(f"写入 Excel 失败：{exc}")
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass
